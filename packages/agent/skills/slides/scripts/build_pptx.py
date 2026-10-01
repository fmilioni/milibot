#!/usr/bin/env python3
"""Builds a 16:9 PowerPoint deck from slide images (one per slide, in file-name order) and optional notes.

Run with: uv run --with python-pptx python build_pptx.py --images <folder> --out deck.pptx [--notes notes.json]
notes.json maps an image name (with or without extension) to that slide's speaker notes.
"""

import argparse
import json
import pathlib
import sys

from pptx import Presentation
from pptx.util import Emu

SLIDE_WIDTH = Emu(12192000)  # 13.333 in
SLIDE_HEIGHT = Emu(6858000)  # 7.5 in
BLANK_LAYOUT = 6


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--images", required=True, help="folder with the slide images (png/jpg)")
    parser.add_argument("--out", required=True, help="output .pptx")
    parser.add_argument("--notes", help="JSON file {image name: speaker notes}")
    args = parser.parse_args()

    folder = pathlib.Path(args.images)
    if not folder.is_dir():
        sys.exit(f"not a folder: {folder}")
    images = sorted(f for f in folder.iterdir() if f.suffix.lower() in (".png", ".jpg", ".jpeg"))
    if not images:
        sys.exit(f"no png/jpg images in {folder}")
    notes = json.loads(pathlib.Path(args.notes).read_text(encoding="utf-8")) if args.notes else {}

    deck = Presentation()
    deck.slide_width = SLIDE_WIDTH
    deck.slide_height = SLIDE_HEIGHT
    for image in images:
        slide = deck.slides.add_slide(deck.slide_layouts[BLANK_LAYOUT])
        slide.shapes.add_picture(str(image), 0, 0, width=SLIDE_WIDTH, height=SLIDE_HEIGHT)
        text = notes.get(image.name) or notes.get(image.stem)
        if text:
            slide.notes_slide.notes_text_frame.text = str(text)

    out = pathlib.Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    deck.save(out)
    print(f"{out}: {len(images)} slides")


if __name__ == "__main__":
    main()
